import random
import string
import openpyxl
from io import BytesIO
from django.utils.text import slugify
from accounts.models import User
from django.core.files.base import ContentFile

class StudentManagementService:
    @staticmethod
    def generate_username(full_name, dob, domain):
        """
        Generates an ID like john.doe.15052010@thaagam.org
        dob should be in DDMMYYYY format
        """
        slug_name = slugify(full_name).replace('-', '.')
        # Remove any leading/trailing dots just in case
        slug_name = slug_name.strip('.')
        # Clean domain
        clean_domain = domain.strip().lower().replace('https://', '').replace('http://', '').split('/')[0]
        
        email = f"{slug_name}.{dob}@{clean_domain}"
        
        # Handle duplicates by adding a suffix if needed
        base_email = email
        counter = 1
        name_part, domain_part = base_email.split('@')
        while User.objects.filter(email__iexact=email).exists():
            email = f"{name_part}.{counter}@{domain_part}"
            counter += 1
            
        return email

    @staticmethod
    def generate_random_password(length=8):
        characters = string.ascii_letters + string.digits
        return ''.join(random.choice(characters) for _ in range(length))

    @staticmethod
    def create_student(full_name, dob, domain, school, teacher=None):
        from accounts.models import StudentCredentialLog
        email = StudentManagementService.generate_username(full_name, dob, domain)
        password = StudentManagementService.generate_random_password()
        
        # Split name into first and last
        parts = full_name.split(' ', 1)
        first_name = parts[0]
        last_name = parts[1] if len(parts) > 1 else ''
        
        user = User.objects.create_user(
            email=email,
            username=full_name,
            first_name=first_name,
            last_name=last_name,
            password=password,
            role='student',
            school=school,
            email_verified=True,
            is_active=True
        )

        # Log the credentials
        StudentCredentialLog.objects.create(
            school=school,
            student_name=full_name,
            generated_id=email,
            plain_password=password,
            teacher=teacher,
            action_type='create'
        )
        
        return user, password

    @staticmethod
    def process_bulk_excel(file_obj, domain, school, teacher=None):
        wb = openpyxl.load_workbook(file_obj)
        sheet = wb.active
        
        results = []
        # Expecting headers: Full Name, DOB (DDMMYYYY)
        # Skip header if present (assuming row 1 is header)
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            full_name = str(row[0]).strip() if row[0] else None
            dob = str(row[1]).strip() if row[1] else None
            
            if not full_name or not dob:
                continue
                
            try:
                user, password = StudentManagementService.create_student(full_name, dob, domain, school, teacher=teacher)
                results.append({
                    'name': full_name,
                    'dob': dob,
                    'email': user.email,
                    'password': password,
                    'status': 'Success'
                })
            except Exception as e:
                results.append({
                    'name': full_name,
                    'dob': dob,
                    'email': 'N/A',
                    'password': 'N/A',
                    'status': f'Error: {str(e)}'
                })
        
        # Create output workbook
        out_wb = openpyxl.Workbook()
        out_sheet = out_wb.active
        out_sheet.title = "Generated Credentials"
        
        headers = ["Full Name", "DOB", "System ID (Email)", "Password", "Status"]
        out_sheet.append(headers)
        
        for res in results:
            out_sheet.append([res['name'], res['dob'], res['email'], res['password'], res['status']])
            
        output = BytesIO()
        out_wb.save(output)
        output.seek(0)
        return output.getvalue()
