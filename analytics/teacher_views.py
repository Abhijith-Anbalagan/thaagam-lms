from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from accounts.decorators import role_required

from analytics.services.analytics_services import TeacherAnalyticsService


@login_required
@role_required('teacher', 'management', 'school_admin', 'super_admin')
def teacher_analytics(request):
    """Teacher analytics dashboard (Global)."""
    service = TeacherAnalyticsService(request.user)

    context = {
        'class_performance': service.get_class_performance(),
        'student_performance': service.get_student_wise_performance(),
        'assignment_analytics': service.get_assignment_analytics(),
        'weak_students': service.identify_weak_students(),
    }

    return render(request, 'teacher/analytics.html', context)


@login_required
@role_required('teacher')
def classroom_analytics(request, classroom_id):
    """Classroom-specific analytics dashboard."""
    service = TeacherAnalyticsService(request.user, classroom_id=classroom_id)
    metrics = service.get_classroom_performance_metrics(classroom_id)

    if not metrics:
        from django.shortcuts import get_object_or_404
        from classrooms.models import Classroom
        get_object_or_404(Classroom, id=classroom_id, teacher=request.user) # trigger 404 if not found/authorized

    context = {
        'classroom': metrics['classroom'],
        'student_performance': metrics['students_performance'],
        'assignment_analytics': metrics['assignments_analytics'],
        'weak_students': metrics['weak_students'],
        'active_tab': 'analytics',
    }

    return render(request, 'teacher/classroom_analytics.html', context)


@login_required
@role_required('teacher')
def export_overall_analytics(request):
    """Export all classroom analytics for the teacher to Excel."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from django.http import HttpResponse
    from datetime import datetime

    service = TeacherAnalyticsService(request.user)
    
    # Get Data
    class_performance = service.get_class_performance()
    student_performance = service.get_student_wise_performance()
    assignment_analytics = service.get_assignment_analytics()
    
    # Create Workbook
    wb = openpyxl.Workbook()
    
    # Sheet 1: Class-wise Performance
    ws1 = wb.active
    ws1.title = "Class Performance"
    headers1 = ["Classroom Name", "Class Code", "Total Students", "Assignments", "Avg Performance %"]
    ws1.append(headers1)
    
    header_fill = PatternFill(start_color="0F6E56", end_color="0F6E56", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws1[1]:
        cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center")

    for c in class_performance:
        ws1.append([
            c.get('name', 'N/A'),
            c.get('code', 'N/A'),
            c.get('total_students_count', 0),
            c.get('assignment_count', 0),
            round(c.get('avg_score', 0) or 0, 2)
        ])

    # Sheet 2: Student Performance
    ws2 = wb.create_sheet(title="Student Performance")
    headers2 = ["Student Name", "Email", "Enrollments", "Assignments", "Submissions", "Avg Score %"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center")

    for s in student_performance:
        ws2.append([
            s.get('username', 'N/A'),
            s.get('email', 'N/A'),
            s.get('classroom_count', 0),
            s.get('assignment_count', 0),
            s.get('submission_count', 0),
            round(s.get('avg_score', 0) or 0, 2)
        ])

    # Sheet 3: Assignment History
    ws3 = wb.create_sheet(title="Assignments")
    headers3 = ["Assignment Title", "Classroom", "Total Students", "Submissions", "Late", "Avg Score %"]
    ws3.append(headers3)
    for cell in ws3[1]:
        cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center")

    for a in assignment_analytics:
        ws3.append([
            a.get('title', 'N/A'),
            a.get('classroom__name', 'N/A'),
            a.get('total_students', 0),
            a.get('submissions_count', 0),
            a.get('late_submissions', 0),
            round(a.get('avg_score', 0) or 0, 2)
        ])

    # Adjust Column Widths
    for ws in [ws1, ws2, ws3]:
        for column_cells in ws.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = length + 2

    # Prepare Response
    filename = f"Overall_Analytics_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response


@login_required
@role_required('teacher')
def export_classroom_analytics(request, classroom_id):
    """Export classroom analytics to Excel."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    from django.http import HttpResponse
    from classrooms.models import Classroom
    from django.shortcuts import get_object_or_404
    from datetime import datetime

    classroom = get_object_or_404(Classroom, id=classroom_id, teacher=request.user)
    service = TeacherAnalyticsService(request.user, classroom_id=classroom_id)
    
    # Get Data
    student_performance = service.get_student_wise_performance(classroom_id)
    assignment_analytics = service.get_assignment_analytics(classroom_id)
    
    # Create Workbook
    wb = openpyxl.Workbook()
    
    # Sheet 1: Student Performance
    ws1 = wb.active
    ws1.title = "Student Performance"
    
    headers = ["Student Name", "Email", "Enrollments", "Assignments", "Submissions", "Avg Score %", "Completed"]
    ws1.append(headers)
    
    # Style Headers
    header_fill = PatternFill(start_color="0F6E56", end_color="0F6E56", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws1[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for s in student_performance:
        ws1.append([
            s.get('username', 'N/A'),
            s.get('email', 'N/A'),
            s.get('classroom_count', 0),
            s.get('assignment_count', 0),
            s.get('submission_count', 0),
            round(s.get('avg_score', 0) or 0, 2),
            s.get('completed_count', 0)
        ])

    # Sheet 2: Assignment Analytics
    ws2 = wb.create_sheet(title="Assignment Analytics")
    headers2 = ["Assignment Title", "Due Date", "Total Students", "Submissions", "Late", "Avg Score %", "Submission Rate %"]
    ws2.append(headers2)
    
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for a in assignment_analytics:
        due_dt = a.get('due_date')
        due_date_str = due_dt.strftime("%Y-%m-%d %H:%M") if due_dt else "N/A"
        ws2.append([
            a.get('title', 'N/A'),
            due_date_str,
            a.get('total_students', 0),
            a.get('submissions_count', 0),
            a.get('late_submissions', 0),
            round(a.get('avg_score', 0) or 0, 2),
            a.get('submission_rate', 0)
        ])

    # Adjust Column Widths
    for ws in [ws1, ws2]:
        for column_cells in ws.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = length + 2

    # Prepare Response
    filename = f"Analytics_{classroom.name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    
    return response